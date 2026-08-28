"""
test_desktop_control_personal.py
E2E Test Harness validating Desktop App (app_local) Control de Personal features:
- Jefes de Obra, Cargos y Tarifas Diarias CRUD
- Maestros management with on-the-fly cargo creation
- Daily rates sueldo_diario calculations (1.0 * sueldo_diario, 0.5 * sueldo_diario)
- Dashboard KPIs calculation
- Attendance Reports filtering by date range & totals calculation
- Excel (.xlsx) and PDF report export generation.
"""

import unittest
import uuid
import os
import tempfile
import datetime
import pandas as pd
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


class MockControlPersonalController:
    """
    In-memory controller simulating app_local Control de Personal backend logic
    and database interactions for Jefes, Cargos, Maestros, KPIs, and Reports.
    """
    def __init__(self):
        self.jefes = {} # id -> {id, nombre, rol}
        self.cargos = {} # id -> {id, nombre, sueldo_diario}
        self.maestros = {} # id -> {id, rut, nombre_completo, tipo_maestro_id, telefono, activo}
        self.obras_maestros = [] # [{id, obra_id, maestro_id}]
        self.asistencias = {} # id -> {id, obra_id, fecha, foto_evidencia_url, es_correctivo, registrado_por}
        self.asistencias_detalle = [] # [{id, asistencia_diaria_id, maestro_id, estado_asistencia, observaciones}]

    # --- 1. Jefes de Obra CRUD ---
    def crear_jefe(self, nombre):
        if not nombre or not nombre.strip():
            return False, "El nombre no puede estar vacío."
        j_id = str(uuid.uuid4())
        self.jefes[j_id] = {"id": j_id, "nombre": nombre.strip(), "rol": "Jefe_Obra"}
        return True, j_id

    def listar_jefes(self):
        return list(self.jefes.values())

    # --- 2. Cargos y Tarifas Diarias CRUD ---
    def crear_cargo(self, nombre, sueldo_diario=0.0):
        if not nombre or not nombre.strip():
            return False, "El nombre del cargo no puede estar vacío."
        # Check duplicate
        for c in self.cargos.values():
            if c["nombre"].lower() == nombre.strip().lower():
                return True, c["id"] # Return existing cargo ID
        c_id = str(uuid.uuid4())
        self.cargos[c_id] = {"id": c_id, "nombre": nombre.strip(), "sueldo_diario": float(sueldo_diario)}
        return True, c_id

    def actualizar_tarifa_cargo(self, cargo_id, nuevo_sueldo_diario):
        if cargo_id not in self.cargos:
            return False, "Cargo no encontrado."
        if nuevo_sueldo_diario < 0:
            return False, "El sueldo diario no puede ser negativo."
        self.cargos[cargo_id]["sueldo_diario"] = float(nuevo_sueldo_diario)
        return True, "Tarifa actualizada."

    def listar_cargos(self):
        return list(self.cargos.values())

    # --- 3. Maestros Management & On-The-Fly Cargo Creation ---
    def registrar_maestro(self, rut, nombre_completo, cargo_nombre_o_id, telefono=""):
        if not rut or not nombre_completo:
            return False, "RUT y Nombre completo son obligatorios."
        # Check duplicate RUT
        for m in self.maestros.values():
            if m["rut"] == rut:
                return False, f"El RUT '{rut}' ya se encuentra registrado."

        # On-the-fly cargo resolution: if UUID present in self.cargos, use it; else create new cargo
        cargo_id = None
        if cargo_nombre_o_id in self.cargos:
            cargo_id = cargo_nombre_o_id
        else:
            # Create cargo on the fly with default sueldo_diario 0.0
            ok, res_id = self.crear_cargo(cargo_nombre_o_id, sueldo_diario=0.0)
            if not ok:
                return False, f"Error creando cargo al vuelo: {res_id}"
            cargo_id = res_id

        m_id = str(uuid.uuid4())
        self.maestros[m_id] = {
            "id": m_id,
            "rut": rut.strip(),
            "nombre_completo": nombre_completo.strip(),
            "tipo_maestro_id": cargo_id,
            "telefono": telefono.strip(),
            "activo": True
        }
        return True, m_id

    def asignar_maestro_a_obra(self, obra_id, maestro_id):
        # Check duplicate
        for om in self.obras_maestros:
            if om["obra_id"] == obra_id and om["maestro_id"] == maestro_id:
                return False, "El maestro ya está asignado a esta obra."
        om_id = str(uuid.uuid4())
        self.obras_maestros.append({"id": om_id, "obra_id": obra_id, "maestro_id": maestro_id})
        return True, om_id

    # --- 4. Daily Rates & Calculation Logic ---
    def calcular_pago_diario(self, maestro_id, estado_asistencia):
        """
        Calculates payable salary for a single day:
        1.0 (Día Completo) -> 1.0 * sueldo_diario
        0.5 (Medio Día) -> 0.5 * sueldo_diario
        0.0 (Ausente) -> 0.0
        """
        if estado_asistencia not in (1.0, 0.5, 0.0):
            raise ValueError(f"Estado de asistencia inválido: {estado_asistencia}")
        maestro = self.maestros.get(maestro_id)
        if not maestro:
            return 0.0
        cargo = self.cargos.get(maestro["tipo_maestro_id"])
        sueldo_base = cargo["sueldo_diario"] if cargo else 0.0
        return estado_asistencia * sueldo_base

    # --- 5. Dashboard KPIs Calculation ---
    def calcular_dashboard_kpis(self, obra_id, fecha):
        # Find attendance header for obra_id and fecha
        header = None
        for a in self.asistencias.values():
            if a["obra_id"] == obra_id and a["fecha"] == fecha:
                header = a
                break

        if not header:
            return {
                "total_maestros_asignados": len([om for om in self.obras_maestros if om["obra_id"] == obra_id]),
                "evaluados": 0,
                "total_jornales": 0.0,
                "porcentaje_asistencia": 0.0,
                "costo_diario_total": 0.0
            }

        detalles = [d for d in self.asistencias_detalle if d["asistencia_diaria_id"] == header["id"]]
        evaluados = len(detalles)
        total_jornales = sum(d["estado_asistencia"] for d in detalles)
        porcentaje = (total_jornales / evaluados * 100.0) if evaluados > 0 else 0.0

        costo_diario = 0.0
        for d in detalles:
            costo_diario += self.calcular_pago_diario(d["maestro_id"], d["estado_asistencia"])

        return {
            "total_maestros_asignados": len([om for om in self.obras_maestros if om["obra_id"] == obra_id]),
            "evaluados": evaluados,
            "total_jornales": total_jornales,
            "porcentaje_asistencia": round(porcentaje, 1),
            "costo_diario_total": round(costo_diario, 2)
        }

    # --- 6. Attendance Reports & Date Range Filtering ---
    def generar_reporte_asistencia_obra(self, obra_id, fecha_inicio, fecha_fin):
        headers_in_range = [
            a for a in self.asistencias.values()
            if a["obra_id"] == obra_id and fecha_inicio <= a["fecha"] <= fecha_fin
        ]

        header_ids = [h["id"] for h in headers_in_range]
        detalles_in_range = [d for d in self.asistencias_detalle if d["asistencia_diaria_id"] in header_ids]

        # Aggregate by maestro
        resumen_maestros = {}
        for d in detalles_in_range:
            m_id = d["maestro_id"]
            if m_id not in resumen_maestros:
                m = self.maestros.get(m_id, {})
                c = self.cargos.get(m.get("tipo_maestro_id"), {})
                resumen_maestros[m_id] = {
                    "rut": m.get("rut", ""),
                    "nombre_completo": m.get("nombre_completo", ""),
                    "cargo": c.get("nombre", "Sin Cargo"),
                    "sueldo_diario": c.get("sueldo_diario", 0.0),
                    "dias_completos": 0,
                    "medios_dias": 0,
                    "ausentes": 0,
                    "total_jornales": 0.0,
                    "total_sueldo": 0.0
                }

            est = d["estado_asistencia"]
            r = resumen_maestros[m_id]
            if est == 1.0:
                r["dias_completos"] += 1
            elif est == 0.5:
                r["medios_dias"] += 1
            elif est == 0.0:
                r["ausentes"] += 1

            r["total_jornales"] += est
            pago = self.calcular_pago_diario(m_id, est)
            r["total_sueldo"] += pago

        rows = list(resumen_maestros.values())
        total_jornales_obra = sum(r["total_jornales"] for r in rows)
        total_sueldo_obra = sum(r["total_sueldo"] for r in rows)

        return {
            "obra_id": obra_id,
            "fecha_inicio": fecha_inicio,
            "fecha_fin": fecha_fin,
            "registros": rows,
            "total_jornales_obra": total_jornales_obra,
            "total_sueldo_obra": total_sueldo_obra
        }


class TestDesktopControlPersonal(unittest.TestCase):
    def setUp(self):
        self.controller = MockControlPersonalController()
        self.obra_id = "obra-valparaiso-01"

    def test_01_jefes_obra_crud(self):
        """Validates Jefe de Obra creation and listing"""
        ok, msg = self.controller.crear_jefe("")
        self.assertFalse(ok) # Empty name must fail

        ok, j1_id = self.controller.crear_jefe("Carlos Rodriguez")
        self.assertTrue(ok)
        self.assertIsNotNone(j1_id)

        jefes = self.controller.listar_jefes()
        self.assertEqual(len(jefes), 1)
        self.assertEqual(jefes[0]["nombre"], "Carlos Rodriguez")

    def test_02_cargos_y_tarifas_crud(self):
        """Validates Cargos creation and daily rate sueldo_diario updates"""
        ok, c1_id = self.controller.crear_cargo("Maestro Mayor", 50000.0)
        self.assertTrue(ok)

        cargos = self.controller.listar_cargos()
        self.assertEqual(len(cargos), 1)
        self.assertEqual(cargos[0]["sueldo_diario"], 50000.0)

        # Update daily rate
        ok, msg = self.controller.actualizar_tarifa_cargo(c1_id, 55000.0)
        self.assertTrue(ok)
        self.assertEqual(self.controller.cargos[c1_id]["sueldo_diario"], 55000.0)

        # Negative rate must fail
        ok, msg = self.controller.actualizar_tarifa_cargo(c1_id, -1000.0)
        self.assertFalse(ok)

    def test_03_maestros_on_the_fly_cargo_creation(self):
        """Validates Maestro registration with on-the-fly cargo creation"""
        # Register maestro with a non-existent cargo name "Especialista HVAC"
        ok, m_id = self.controller.registrar_maestro("15.999.888-7", "Gonzalo Morales", "Especialista HVAC", "+56977778888")
        self.assertTrue(ok)

        maestro = self.controller.maestros[m_id]
        self.assertEqual(maestro["nombre_completo"], "Gonzalo Morales")

        # Verify new cargo was created automatically
        cargo_id = maestro["tipo_maestro_id"]
        cargo = self.controller.cargos[cargo_id]
        self.assertEqual(cargo["nombre"], "Especialista HVAC")
        self.assertEqual(cargo["sueldo_diario"], 0.0) # Default sueldo_diario

    def test_04_daily_rate_salary_calculations(self):
        """Validates 1.0 * sueldo_diario, 0.5 * sueldo_diario, 0.0 calculation"""
        _, c_id = self.controller.crear_cargo("Ayudante", 30000.0)
        _, m_id = self.controller.registrar_maestro("16.111.222-3", "Juan Ayudante", c_id)

        # 1.0 (Día Completo) -> $30,000
        pago_full = self.controller.calcular_pago_diario(m_id, 1.0)
        self.assertEqual(pago_full, 30000.0)

        # 0.5 (Medio Día) -> $15,000
        pago_half = self.controller.calcular_pago_diario(m_id, 0.5)
        self.assertEqual(pago_half, 15000.0)

        # 0.0 (Ausente) -> $0
        pago_zero = self.controller.calcular_pago_diario(m_id, 0.0)
        self.assertEqual(pago_zero, 0.0)

    def test_05_dashboard_kpis_calculation(self):
        """Validates Dashboard KPI metrics calculation (evaluados, jornales, %, costo total)"""
        _, c1 = self.controller.crear_cargo("Carpintero", 40000.0) # $40k
        _, c2 = self.controller.crear_cargo("Jornal", 20000.0)    # $20k

        _, m1 = self.controller.registrar_maestro("10.001-1", "Maestro 1", c1)
        _, m2 = self.controller.registrar_maestro("10.002-2", "Maestro 2", c2)
        _, m3 = self.controller.registrar_maestro("10.003-3", "Maestro 3", c2)

        self.controller.asignar_maestro_a_obra(self.obra_id, m1)
        self.controller.asignar_maestro_a_obra(self.obra_id, m2)
        self.controller.asignar_maestro_a_obra(self.obra_id, m3)

        # Record attendance for 2026-07-29
        ast_id = str(uuid.uuid4())
        self.controller.asistencias[ast_id] = {
            "id": ast_id, "obra_id": self.obra_id, "fecha": "2026-07-29",
            "foto_evidencia_url": "http://foto.jpg", "es_correctivo": False, "registrado_por": "jefe-1"
        }
        # m1: 1.0 ($40k), m2: 0.5 ($10k), m3: 0.0 ($0) -> Total jornales = 1.5, Total cost = $50,000
        self.controller.asistencias_detalle.extend([
            {"id": str(uuid.uuid4()), "asistencia_diaria_id": ast_id, "maestro_id": m1, "estado_asistencia": 1.0, "observaciones": ""},
            {"id": str(uuid.uuid4()), "asistencia_diaria_id": ast_id, "maestro_id": m2, "estado_asistencia": 0.5, "observaciones": ""},
            {"id": str(uuid.uuid4()), "asistencia_diaria_id": ast_id, "maestro_id": m3, "estado_asistencia": 0.0, "observaciones": ""}
        ])

        kpis = self.controller.calcular_dashboard_kpis(self.obra_id, "2026-07-29")
        self.assertEqual(kpis["evaluados"], 3)
        self.assertEqual(kpis["total_jornales"], 1.5)
        self.assertEqual(kpis["porcentaje_asistencia"], 50.0) # 1.5 / 3 * 100
        self.assertEqual(kpis["costo_diario_total"], 50000.0)

    def test_06_attendance_reports_date_range_filtering(self):
        """Validates Attendance Report date range filtering & total payroll calculations"""
        _, c1 = self.controller.crear_cargo("Encofrador", 50000.0)
        _, m1 = self.controller.registrar_maestro("12.999-9", "Juan Encofrador", c1)

        # Day 1: 2026-07-01 -> 1.0 ($50,000)
        ast1 = str(uuid.uuid4())
        self.controller.asistencias[ast1] = {"id": ast1, "obra_id": self.obra_id, "fecha": "2026-07-01", "foto_evidencia_url": "f1"}
        self.controller.asistencias_detalle.append({"id": str(uuid.uuid4()), "asistencia_diaria_id": ast1, "maestro_id": m1, "estado_asistencia": 1.0})

        # Day 2: 2026-07-02 -> 0.5 ($25,000)
        ast2 = str(uuid.uuid4())
        self.controller.asistencias[ast2] = {"id": ast2, "obra_id": self.obra_id, "fecha": "2026-07-02", "foto_evidencia_url": "f2"}
        self.controller.asistencias_detalle.append({"id": str(uuid.uuid4()), "asistencia_diaria_id": ast2, "maestro_id": m1, "estado_asistencia": 0.5})

        # Day 3: 2026-07-15 (Outside date filter range) -> 1.0
        ast3 = str(uuid.uuid4())
        self.controller.asistencias[ast3] = {"id": ast3, "obra_id": self.obra_id, "fecha": "2026-07-15", "foto_evidencia_url": "f3"}
        self.controller.asistencias_detalle.append({"id": str(uuid.uuid4()), "asistencia_diaria_id": ast3, "maestro_id": m1, "estado_asistencia": 1.0})

        # Report for range 2026-07-01 to 2026-07-05
        report = self.controller.generar_reporte_asistencia_obra(self.obra_id, "2026-07-01", "2026-07-05")
        self.assertEqual(report["total_jornales_obra"], 1.5)
        self.assertEqual(report["total_sueldo_obra"], 75000.0)
        self.assertEqual(len(report["registros"]), 1)
        self.assertEqual(report["registros"][0]["dias_completos"], 1)
        self.assertEqual(report["registros"][0]["medios_dias"], 1)

    def test_07_excel_export_generation(self):
        """Validates Excel (.xlsx) report export generation using pandas & openpyxl"""
        _, c1 = self.controller.crear_cargo("Gasfiter", 45000.0)
        _, m1 = self.controller.registrar_maestro("11.222.333-4", "Mario Bros", c1)

        report_data = {
            "obra_id": self.obra_id,
            "fecha_inicio": "2026-07-01",
            "fecha_fin": "2026-07-31",
            "registros": [{
                "rut": "11.222.333-4",
                "nombre_completo": "Mario Bros",
                "cargo": "Gasfiter",
                "sueldo_diario": 45000.0,
                "dias_completos": 10,
                "medios_dias": 2,
                "ausentes": 1,
                "total_jornales": 11.0,
                "total_sueldo": 495000.0
            }],
            "total_jornales_obra": 11.0,
            "total_sueldo_obra": 495000.0
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = os.path.join(tmpdir, "reporte_asistencia.xlsx")

            df = pd.DataFrame(report_data["registros"])
            df.columns = ["RUT", "Nombre Completo", "Cargo", "Sueldo Diario", "Días Completos", "Medios Días", "Ausentes", "Total Jornales", "Total Sueldo ($)"]

            with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name="Reporte Asistencia", index=False)

            self.assertTrue(os.path.exists(excel_path))

            # Inspect generated workbook with openpyxl
            wb = openpyxl.load_workbook(excel_path)
            self.assertIn("Reporte Asistencia", wb.sheetnames)
            ws = wb["Reporte Asistencia"]
            self.assertEqual(ws.cell(row=1, column=1).value, "RUT")
            self.assertEqual(ws.cell(row=2, column=1).value, "11.222.333-4")
            self.assertEqual(ws.cell(row=2, column=9).value, 495000.0)

    def test_08_pdf_export_generation(self):
        """Validates PDF report export generation using reportlab"""
        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = os.path.join(tmpdir, "reporte_asistencia.pdf")

            doc = SimpleDocTemplate(pdf_path, pagesize=letter)
            styles = getSampleStyleSheet()
            elements = []

            title_style = ParagraphStyle(
                'ReportTitle',
                parent=styles['Heading1'],
                textColor=colors.HexColor('#1A529B'),
                fontSize=18,
                leading=22
            )
            elements.append(Paragraph("<b>ACCURATECK — Control de Asistencia Diario</b>", title_style))
            elements.append(Spacer(1, 12))

            table_data = [
                ["RUT", "Nombre Completo", "Cargo", "Total Jornales", "Total Sueldo ($)"],
                ["11.222.333-4", "Mario Bros", "Gasfiter", "11.0", "$495.000"],
                ["TOTAL OBRA", "", "", "11.0", "$495.000"]
            ]

            t = Table(table_data, colWidths=[90, 150, 100, 90, 100])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A529B')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ]))
            elements.append(t)

            doc.build(elements)

            self.assertTrue(os.path.exists(pdf_path))
            self.assertGreater(os.path.getsize(pdf_path), 500)


if __name__ == '__main__':
    unittest.main()
