from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel, QComboBox,
    QDateEdit, QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView,
    QMessageBox, QFileDialog, QFrame, QSizePolicy, QProgressBar
)
from PyQt6.QtCore import Qt, QDate
import datetime

from src.core.workers import FetchObrasWorker, FetchAsistenciaReporteWorker


class ReportesAsistenciaContent(QWidget):
    def __init__(self, supabase_client=None):
        super().__init__()
        self.supabase_client = supabase_client
        self._report_data = []
        self._report_summary = {}

        layout = QVBoxLayout(self)
        layout.setContentsMargins(30, 30, 30, 30)
        layout.setSpacing(16)

        # ── Título ──
        title = QLabel("Reportes de Asistencia por Obra")
        title.setObjectName("TitleLabel")
        layout.addWidget(title)

        # ── Panel de filtros ──
        filter_frame = QFrame()
        filter_frame.setStyleSheet(
            "QFrame { background: #FAF8F5; color: #1F2421; border: 1px solid #D8D5CD; border-radius: 10px; }"
        )
        filter_layout = QVBoxLayout(filter_frame)
        filter_layout.setContentsMargins(16, 14, 16, 14)
        filter_layout.setSpacing(10)

        filtros_label = QLabel("Filtros del Reporte")
        filtros_label.setStyleSheet("font-weight: bold; font-size: 14px; color: #262B29;")
        filter_layout.addWidget(filtros_label)

        row1 = QHBoxLayout()
        row1.setSpacing(12)

        # Obra
        obra_wrap = QVBoxLayout()
        obra_label = QLabel("Obra:")
        obra_label.setStyleSheet("font-size: 12px; color: #262B29; font-weight: 600;")
        self.combo_obra = QComboBox()
        self.combo_obra.setMinimumWidth(200)
        obra_wrap.addWidget(obra_label)
        obra_wrap.addWidget(self.combo_obra)
        row1.addLayout(obra_wrap)

        # Fecha Desde
        desde_wrap = QVBoxLayout()
        desde_label = QLabel("Fecha Desde:")
        desde_label.setStyleSheet("font-size: 12px; color: #262B29; font-weight: 600;")
        self.date_desde = QDateEdit()
        self.date_desde.setCalendarPopup(True)
        self.date_desde.setDate(QDate.currentDate().addDays(-30))
        self.date_desde.setDisplayFormat("dd/MM/yyyy")
        desde_wrap.addWidget(desde_label)
        desde_wrap.addWidget(self.date_desde)
        row1.addLayout(desde_wrap)

        # Fecha Hasta
        hasta_wrap = QVBoxLayout()
        hasta_label = QLabel("Fecha Hasta:")
        hasta_label.setStyleSheet("font-size: 12px; color: #262B29; font-weight: 600;")
        self.date_hasta = QDateEdit()
        self.date_hasta.setCalendarPopup(True)
        self.date_hasta.setDate(QDate.currentDate())
        self.date_hasta.setDisplayFormat("dd/MM/yyyy")
        hasta_wrap.addWidget(hasta_label)
        hasta_wrap.addWidget(self.date_hasta)
        row1.addLayout(hasta_wrap)

        row1.addStretch()

        # Botones de rango rápido
        quick_wrap = QVBoxLayout()
        quick_label = QLabel("Período Rápido:")
        quick_label.setStyleSheet("font-size: 12px; color: #262B29; font-weight: 600;")
        quick_btns = QHBoxLayout()
        quick_btns.setSpacing(6)

        for label, days in [("Hoy", 0), ("Esta Semana", 6), ("Este Mes", 29)]:
            btn = QPushButton(label)
            btn.setObjectName("SecondaryButton")
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
            btn.setFixedHeight(30)
            btn.clicked.connect(lambda checked, d=days: self._set_quick_range(d))
            quick_btns.addWidget(btn)

        quick_wrap.addWidget(quick_label)
        quick_wrap.addLayout(quick_btns)
        row1.addLayout(quick_wrap)

        filter_layout.addLayout(row1)

        btn_generar = QPushButton("Generar Reporte")
        btn_generar.setObjectName("PrimaryButton")
        btn_generar.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_generar.setFixedHeight(40)
        btn_generar.clicked.connect(self._generar_reporte)
        filter_layout.addWidget(btn_generar, alignment=Qt.AlignmentFlag.AlignLeft)

        layout.addWidget(filter_frame)

        # ── Barra de progreso ──
        self.progress = QProgressBar()
        self.progress.setVisible(False)
        self.progress.setRange(0, 0)
        self.progress.setFixedHeight(4)
        layout.addWidget(self.progress)

        # ── Tabla de resultados ──
        self.table = QTableWidget()
        self.table.verticalHeader().setDefaultSectionSize(44)
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().setMinimumSectionSize(110)
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels([
            "Maestro", "Tipo/Cargo", "Sueldo Diario",
            "Días Completos", "Medios Días", "Jornales Totales", "Pago Total"
        ])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        for col in range(1, 7):
            self.table.horizontalHeader().setSectionResizeMode(col, QHeaderView.ResizeMode.ResizeToContents)
        self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        layout.addWidget(self.table)

        # ── Fila de totales ──
        self.lbl_total = QLabel("")
        self.lbl_total.setStyleSheet(
            "font-weight: bold; font-size: 14px; color: #1F2421; "
            "background: #F2EFE9; border-radius: 6px; padding: 8px 12px;"
        )
        self.lbl_total.setVisible(False)
        layout.addWidget(self.lbl_total)

        # ── Botones de exportación ──
        export_row = QHBoxLayout()
        export_row.setSpacing(10)

        self.btn_excel = QPushButton("Exportar a Excel (.xlsx)")
        self.btn_excel.setObjectName("SecondaryButton")
        self.btn_excel.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_excel.setEnabled(False)
        self.btn_excel.clicked.connect(self._exportar_excel)

        self.btn_pdf = QPushButton("Exportar a PDF")
        self.btn_pdf.setObjectName("SecondaryButton")
        self.btn_pdf.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_pdf.setEnabled(False)
        self.btn_pdf.clicked.connect(self._exportar_pdf)

        # Verificar si reportlab está disponible
        try:
            import reportlab
            self._reportlab_available = True
        except ImportError:
            self._reportlab_available = False
            self.btn_pdf.setEnabled(False)
            self.btn_pdf.setToolTip("Instala reportlab para habilitar exportación PDF")

        export_row.addWidget(self.btn_excel)
        export_row.addWidget(self.btn_pdf)
        export_row.addStretch()
        layout.addLayout(export_row)

        if self.supabase_client:
            self._cargar_obras()

    def _cargar_obras(self):
        self._w_obras = FetchObrasWorker(self.supabase_client)
        self._w_obras.success.connect(self._on_obras_loaded)
        self._w_obras.start()

    def _on_obras_loaded(self, obras):
        self.combo_obra.clear()
        for o in obras:
            self.combo_obra.addItem(o["nombre"], o["id"])

    def _set_quick_range(self, days):
        today = QDate.currentDate()
        self.date_hasta.setDate(today)
        self.date_desde.setDate(today.addDays(-days))

    def _generar_reporte(self):
        obra_id = self.combo_obra.currentData()
        if not obra_id:
            QMessageBox.warning(self, "Error", "Seleccione una obra.")
            return
        fecha_desde = self.date_desde.date().toString("yyyy-MM-dd")
        fecha_hasta = self.date_hasta.date().toString("yyyy-MM-dd")

        self.progress.setVisible(True)
        self.table.setRowCount(0)
        self.lbl_total.setVisible(False)
        self.btn_excel.setEnabled(False)
        self.btn_pdf.setEnabled(False)

        self._w_rep = FetchAsistenciaReporteWorker(
            self.supabase_client, obra_id, fecha_desde, fecha_hasta
        )
        self._w_rep.success.connect(self._on_reporte_success)
        self._w_rep.error.connect(self._on_reporte_error)
        self._w_rep.start()

    def _on_reporte_success(self, maestros_data, summary):
        self.progress.setVisible(False)
        self._report_data = maestros_data
        self._report_summary = summary

        self.table.setRowCount(len(maestros_data))
        for i, m in enumerate(maestros_data):
            self.table.setItem(i, 0, QTableWidgetItem(m.get("maestro_nombre", "")))
            self.table.setItem(i, 1, QTableWidgetItem(m.get("tipo_nombre", "")))
            sueldo_item = QTableWidgetItem(f"$ {m.get('sueldo_diario', 0):,.2f}")
            sueldo_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(i, 2, sueldo_item)

            dias_item = QTableWidgetItem(str(m.get("dias_completos", 0)))
            dias_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(i, 3, dias_item)

            medios_item = QTableWidgetItem(str(m.get("medios_dias", 0)))
            medios_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(i, 4, medios_item)

            jornales_item = QTableWidgetItem(f"{m.get('jornales_totales', 0):.1f}")
            jornales_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(i, 5, jornales_item)

            pago_item = QTableWidgetItem(f"$ {m.get('pago_total', 0):,.2f}")
            pago_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(i, 6, pago_item)

        total_pago = summary.get("total_pago", 0)
        total_jornales = summary.get("total_jornales", 0)
        obra_nombre = self.combo_obra.currentText()
        self.lbl_total.setText(
            f"Obra: {obra_nombre}  |  Jornales Totales: {total_jornales:.1f}  |  "
            f"Pago Total en el Período: $ {total_pago:,.2f}"
        )
        self.lbl_total.setVisible(True)
        self.btn_excel.setEnabled(bool(maestros_data))
        if self._reportlab_available:
            self.btn_pdf.setEnabled(bool(maestros_data))

    def _on_reporte_error(self, msg):
        self.progress.setVisible(False)
        QMessageBox.critical(self, "Error", msg)

    def _exportar_excel(self):
        if not self._report_data:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            QMessageBox.critical(self, "Error", "Instala 'openpyxl' para exportar a Excel.\npip install openpyxl")
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "Guardar Reporte de Asistencia", 
            f"asistencia_{self.combo_obra.currentText().replace(' ', '_')}_{datetime.date.today()}.xlsx",
            "Excel (*.xlsx)"
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        obra_nombre = self.combo_obra.currentText()
        ws.title = "Asistencia"
        ws["A1"] = f"Reporte de Asistencia — {obra_nombre}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (
            f"Período: {self.date_desde.date().toString('dd/MM/yyyy')} "
            f"al {self.date_hasta.date().toString('dd/MM/yyyy')}"
        )
        ws["A2"].font = Font(size=11, color="64748B")

        headers = ["Maestro", "Tipo/Cargo", "Sueldo Diario", "Días Completos", "Medios Días", "Jornales Totales", "Pago Total"]
        header_fill = PatternFill("solid", fgColor="262B29")
        header_font = Font(bold=True, color="FFFFFF")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, m in enumerate(self._report_data, 5):
            ws.cell(row=row_idx, column=1, value=m.get("maestro_nombre", ""))
            ws.cell(row=row_idx, column=2, value=m.get("tipo_nombre", ""))
            ws.cell(row=row_idx, column=3, value=float(m.get("sueldo_diario", 0)))
            ws.cell(row=row_idx, column=4, value=m.get("dias_completos", 0))
            ws.cell(row=row_idx, column=5, value=m.get("medios_dias", 0))
            ws.cell(row=row_idx, column=6, value=float(m.get("jornales_totales", 0)))
            ws.cell(row=row_idx, column=7, value=float(m.get("pago_total", 0)))

        # Fila de totales
        total_row = len(self._report_data) + 5
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=float(self._report_summary.get("total_jornales", 0))).font = Font(bold=True)
        ws.cell(row=total_row, column=7, value=float(self._report_summary.get("total_pago", 0))).font = Font(bold=True)

        for col in range(1, 8):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].auto_size = True

        wb.save(path)
        QMessageBox.information(self, "Éxito", f"Reporte exportado exitosamente:\n{path}")

    def _exportar_pdf(self):
        if not self._report_data or not self._reportlab_available:
            return
        try:
            from reportlab.lib.pagesizes import letter, landscape
            from reportlab.lib import colors
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
        except ImportError:
            QMessageBox.critical(self, "Error", "Instala 'reportlab' para exportar a PDF.")
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "Guardar Reporte PDF",
            f"asistencia_{self.combo_obra.currentText().replace(' ', '_')}_{datetime.date.today()}.pdf",
            "PDF (*.pdf)"
        )
        if not path:
            return

        doc = SimpleDocTemplate(path, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        elements = []

        obra_nombre = self.combo_obra.currentText()
        elements.append(Paragraph(f"Reporte de Asistencia — {obra_nombre}", styles["Title"]))
        fecha_str = (
            f"Período: {self.date_desde.date().toString('dd/MM/yyyy')} "
            f"al {self.date_hasta.date().toString('dd/MM/yyyy')}"
        )
        elements.append(Paragraph(fecha_str, styles["Normal"]))
        elements.append(Spacer(1, 16))

        data = [["Maestro", "Tipo/Cargo", "Sueldo Diario", "Días Compl.", "Medios Días", "Jornales", "Pago Total"]]
        for m in self._report_data:
            data.append([
                m.get("maestro_nombre", ""),
                m.get("tipo_nombre", ""),
                f"$ {m.get('sueldo_diario', 0):,.2f}",
                str(m.get("dias_completos", 0)),
                str(m.get("medios_dias", 0)),
                f"{m.get('jornales_totales', 0):.1f}",
                f"$ {m.get('pago_total', 0):,.2f}",
            ])
        data.append([
            "TOTAL", "", "",
            "", "",
            f"{self._report_summary.get('total_jornales', 0):.1f}",
            f"$ {self._report_summary.get('total_pago', 0):,.2f}",
        ])

        tabla = Table(data, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3D523C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (1, -1), "LEFT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#FAF8F5")]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#D8D5CD")),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8D5CD")),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("PADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(tabla)

        doc.build(elements)
        QMessageBox.information(self, "Éxito", f"PDF exportado exitosamente:\n{path}")
